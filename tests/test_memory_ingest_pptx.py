"""Knowledge-base ingest of .pptx decks: table and grouped-shape text.

A slide table lives in a graphic frame and a group keeps its members in a shape
tree of its own. Neither has a text frame, so an extractor that only reads
``has_text_frame`` shapes indexes none of that text. Decks are built with
python-pptx in ``tmp_path``; nothing touches the network.
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from agentos.memory.ingest import extract_document_text, ingest_document
from agentos.memory.store import LongTermMemoryStore
from agentos.memory.types import MemorySource


def _pricing_deck() -> bytes:
    """One slide: a title, a 2x2 table, and a text box inside a group."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Quarterly pricing"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Plan"
    table.cell(0, 1).text = "Price"
    table.cell(1, 0).text = "Enterprise"
    table.cell(1, 1).text = "$4,200/seat"
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(4), Inches(3), Inches(1))
    box.text_frame.text = "Renewal deadline is March 3"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _save(tmp_path: Path, prs: Presentation, name: str = "deck.pptx") -> Path:
    path = tmp_path / name
    prs.save(str(path))
    return path


def test_table_rows_and_grouped_text_are_extracted_from_a_file(tmp_path: Path) -> None:
    path = tmp_path / "pricing.pptx"
    path.write_bytes(_pricing_deck())

    assert extract_document_text(path) == (
        "[Slide 1]\n"
        "Quarterly pricing\n"
        "Plan | Price\n"
        "Enterprise | $4,200/seat\n"
        "Renewal deadline is March 3"
    )


def test_uploaded_bytes_take_the_same_path_as_a_file() -> None:
    text = extract_document_text(_pricing_deck(), filename="pricing.pptx")

    assert "Enterprise | $4,200/seat" in text
    assert "Renewal deadline is March 3" in text


def test_text_nested_two_groups_deep_is_extracted(tmp_path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    outer = slide.shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    box = inner.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text_frame.text = "Escalation contact is Priya"

    assert extract_document_text(_save(tmp_path, prs)) == "[Slide 1]\nEscalation contact is Priya"


def test_a_slide_holding_only_a_table_is_not_dropped(tmp_path: Path) -> None:
    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[6])
    table = first.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "EMEA"
    second = prs.slides.add_slide(prs.slide_layouts[5])
    second.shapes.title.text = "Next steps"

    assert extract_document_text(_save(tmp_path, prs)) == (
        "[Slide 1]\nRegion | EMEA\n\n[Slide 2]\nNext steps"
    )


def test_a_merged_table_cell_is_read_once(tmp_path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(2, 3, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Quarterly revenue"
    table.cell(0, 0).merge(table.cell(0, 2))
    table.cell(1, 0).text = "Q1"
    table.cell(1, 1).text = "Q2"
    table.cell(1, 2).text = "Q3"

    assert extract_document_text(_save(tmp_path, prs)) == (
        "[Slide 1]\nQuarterly revenue\nQ1 | Q2 | Q3"
    )


def test_plain_text_frames_read_as_before(tmp_path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Agenda"
    body = slide.placeholders[1].text_frame
    body.text = "Kickoff"
    body.add_paragraph().text = "   "
    body.add_paragraph().text = "Budget review"

    assert extract_document_text(_save(tmp_path, prs)) == (
        "[Slide 1]\nAgenda\nKickoff\nBudget review"
    )


@pytest.mark.asyncio
async def test_knowledge_base_search_finds_table_and_grouped_text(tmp_path: Path) -> None:
    deck = tmp_path / "pricing.pptx"
    deck.write_bytes(_pricing_deck())
    store = LongTermMemoryStore(tmp_path / "memory.db")
    await store.initialize()
    try:
        result = await ingest_document(store, deck, rel_path="knowledge_base/pricing.pptx")
        assert result.status == "indexed"

        # Positive control: the title was indexed before the fix as well.
        title_hits, _ = await store.search(
            "Quarterly", source=MemorySource.knowledge_base, min_score=0.0
        )
        assert [hit.path for hit in title_hits] == ["knowledge_base/pricing.pptx"]

        for query in ("Enterprise", "Renewal deadline"):
            hits, _ = await store.search(query, source=MemorySource.knowledge_base, min_score=0.0)
            assert [hit.path for hit in hits] == ["knowledge_base/pricing.pptx"], query
    finally:
        await store.close()
