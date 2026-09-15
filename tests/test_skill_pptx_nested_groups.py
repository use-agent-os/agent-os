"""pptx skill extract_text: nested group shapes (#1894).

`_slide_text` used to expand exactly one level of grouping, so a text box or a
table sitting inside Group -> Group vanished from the output without a trace.
"""

from __future__ import annotations

import io
import json
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from agentos.skills.bundled.pptx.scripts import extract_text


def _nested_group_deck(tmp_path: Path) -> Path:
    """One slide: text at group depths 1, 2 and 3 plus a table at depth 2."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    top = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(3), Inches(0.5))
    top.text_frame.text = "top level"

    outer = slide.shapes.add_group_shape()
    depth_one = outer.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(0.5))
    depth_one.text_frame.text = "depth one"

    inner = outer.shapes.add_group_shape()
    depth_two = inner.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(0.5))
    depth_two.text_frame.text = "depth two"

    deepest = inner.shapes.add_group_shape()
    depth_three = deepest.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(0.5))
    depth_three.text_frame.text = "depth three"

    # python-pptx only offers add_table on the slide, so build the table there
    # and move its graphicFrame into the depth-two group's shape tree.
    table_shape = slide.shapes.add_table(1, 2, Inches(1), Inches(4), Inches(4), Inches(0.5))
    table_shape.table.cell(0, 0).text_frame.text = "Metric"
    table_shape.table.cell(0, 1).text_frame.text = "Value"
    inner.shapes._spTree.append(table_shape._element)

    path = tmp_path / "nested.pptx"
    prs.save(str(path))
    return path


def test_slide_text_recurses_through_nested_groups(tmp_path: Path) -> None:
    slide = Presentation(str(_nested_group_deck(tmp_path))).slides[0]

    assert extract_text._slide_text(slide) == [
        "top level",
        "depth one",
        "depth two",
        "depth three",
        "Metric | Value",
    ]


def test_cli_json_output_includes_nested_group_text(tmp_path: Path) -> None:
    path = _nested_group_deck(tmp_path)

    stdout = io.StringIO()
    monkeypatch = pytest.MonkeyPatch()
    monkeypatch.setattr(sys, "stdout", stdout)
    try:
        assert extract_text.main([str(path), "--json"]) == 0
    finally:
        monkeypatch.undo()

    (record,) = json.loads(stdout.getvalue())
    assert "depth two" in record["text"]
    assert "depth three" in record["text"]
    assert "Metric | Value" in record["text"]


def test_a_group_whose_children_cannot_be_read_is_skipped_not_fatal() -> None:
    """A malformed group must not take the whole slide's text down with it."""

    class _BrokenGroup:
        shape_type = 6  # MSO_SHAPE_TYPE.GROUP
        has_text_frame = False
        has_table = False

        @property
        def shapes(self) -> object:
            raise AttributeError("no shape tree")

    class _Paragraph:
        text = "still here"

    class _TextFrame:
        paragraphs = [_Paragraph()]

    class _TextShape:
        shape_type = 17  # MSO_SHAPE_TYPE.TEXT_BOX
        has_text_frame = True
        has_table = False
        text_frame = _TextFrame()

    class _Slide:
        shapes = [_BrokenGroup(), _TextShape()]

    assert extract_text._slide_text(_Slide()) == ["still here"]
