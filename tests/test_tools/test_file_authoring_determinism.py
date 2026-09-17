"""Issue #2419: ``create_xlsx`` published different bytes for identical input.

A session deliverable is recognised by the SHA-256 of its bytes, so a tool
whose output carries the wall clock never matches its own previous output and
publishes a fresh artifact every call. openpyxl stamps the current time into
every zip entry and into ``docProps/core.xml`` (``dcterms:created`` and
``dcterms:modified``); ``create_pptx`` already routed its bytes through
``_normalize_zip_timestamps`` and ``create_xlsx`` did not.

``create_pdf_report`` had the same defect from a different writer: reportlab
stamps ``/CreationDate``, ``/ModDate`` and a document ``/ID`` on every save.
Its own ``invariant`` flag exists for exactly this.
"""

from __future__ import annotations

import asyncio
import hashlib
import io
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from agentos.tools.builtin.file_authoring import (
    _normalize_zip_timestamps,
    create_pdf_report,
    create_pptx,
    create_xlsx,
)
from agentos.tools.types import CallerKind, ToolContext, current_tool_context

SHEETS = [{"name": "Summary", "rows": [["metric", "value"], ["requests", 42]]}]
SLIDES = [{"title": "Launch Readiness", "bullets": ["Group reply works", "Artifacts are safe"]}]
DCTERMS = "http://purl.org/dc/terms/"


def _context(tmp_path: Path, *, session: str = "session-1") -> ToolContext:
    workspace = tmp_path / "workspace"
    workspace.mkdir(exist_ok=True)
    return ToolContext(
        caller_kind=CallerKind.CHANNEL,
        workspace_dir=str(workspace),
        artifact_media_root=str(tmp_path / "media"),
        artifact_session_id=session,
        session_key="agent:main:feishu:group:oc_demo",
    )


async def _publish(tmp_path: Path, tool, **kwargs) -> tuple[dict, bytes]:
    """Run *tool* in a fresh context and return its payload and the bytes on disk."""
    import json

    from agentos.artifacts import ArtifactStore

    ctx = _context(tmp_path)
    token = current_tool_context.set(ctx)
    try:
        payload = json.loads(await tool(**kwargs))
    finally:
        current_tool_context.reset(token)
    store = ArtifactStore(ctx.artifact_media_root or "")
    _, path = store.resolve_for_download(
        str(payload["artifact"]["id"]), session_id=str(payload["artifact"]["session_id"])
    )
    return payload, path.read_bytes()


async def _twice_across_the_clock(
    tmp_path: Path, tool, **kwargs
) -> tuple[dict, dict, bytes, bytes]:
    """The issue's shape: two identical calls in one session, more than a second apart."""
    first, first_bytes = await _publish(tmp_path, tool, **kwargs)
    await asyncio.sleep(1.1)
    second, second_bytes = await _publish(tmp_path, tool, **kwargs)
    return first, second, first_bytes, second_bytes


def _core_timestamps(payload: bytes) -> dict[str, str]:
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        root = ElementTree.fromstring(archive.read("docProps/core.xml"))
    return {tag: (root.findtext(f"{{{DCTERMS}}}{tag}") or "") for tag in ("created", "modified")}


def _zip_timestamps(payload: bytes) -> set[tuple[int, ...]]:
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        return {info.date_time for info in archive.infolist()}


# ── the issue: identical input, identical bytes, one artifact ───────────────


@pytest.mark.asyncio
async def test_xlsx_built_twice_is_the_same_bytes(tmp_path: Path) -> None:
    _, _, a, b = await _twice_across_the_clock(tmp_path, create_xlsx, name="m.xlsx", sheets=SHEETS)

    assert a == b
    assert hashlib.sha256(a).hexdigest() == hashlib.sha256(b).hexdigest()


@pytest.mark.asyncio
async def test_xlsx_second_call_is_already_published(tmp_path: Path) -> None:
    """The issue's own reproduction, verbatim in shape."""
    first, second, _, _ = await _twice_across_the_clock(
        tmp_path, create_xlsx, name="metrics.xlsx", sheets=SHEETS
    )

    assert first["status"] == "published"
    assert second["status"] == "already_published"
    assert second["artifact"]["id"] == first["artifact"]["id"]


@pytest.mark.asyncio
async def test_pdf_built_twice_is_the_same_bytes(tmp_path: Path) -> None:
    """The sibling: reportlab's ``/CreationDate``, ``/ModDate`` and ``/ID``."""
    _, _, a, b = await _twice_across_the_clock(
        tmp_path, create_pdf_report, name="r.pdf", title="Report", body="hello"
    )

    assert a == b


@pytest.mark.asyncio
async def test_pdf_second_call_is_already_published(tmp_path: Path) -> None:
    first, second, _, _ = await _twice_across_the_clock(
        tmp_path, create_pdf_report, name="report.pdf", title="Report", body="hello"
    )

    assert first["status"] == "published"
    assert second["status"] == "already_published"
    assert second["artifact"]["id"] == first["artifact"]["id"]


@pytest.mark.asyncio
async def test_pptx_stays_deterministic(tmp_path: Path) -> None:
    """Already true on ``main``; pinned so the shared normaliser cannot regress it."""
    _, _, a, b = await _twice_across_the_clock(tmp_path, create_pptx, name="d.pptx", slides=SLIDES)

    assert a == b


# ── every clock in the package is pinned ────────────────────────────────────


@pytest.mark.asyncio
async def test_xlsx_core_properties_carry_no_wall_clock(tmp_path: Path) -> None:
    _, payload = await _publish(tmp_path, create_xlsx, sheets=SHEETS)

    stamps = _core_timestamps(payload)

    assert stamps == {"created": "1980-01-01T00:00:00Z", "modified": "1980-01-01T00:00:00Z"}


@pytest.mark.asyncio
async def test_xlsx_zip_entries_carry_no_wall_clock(tmp_path: Path) -> None:
    _, payload = await _publish(tmp_path, create_xlsx, sheets=SHEETS)

    assert _zip_timestamps(payload) == {(1980, 1, 1, 0, 0, 0)}


@pytest.mark.asyncio
async def test_pdf_carries_no_wall_clock(tmp_path: Path) -> None:
    _, payload = await _publish(tmp_path, create_pdf_report, title="Report", body="hello")

    for stamp in (rb"/CreationDate \(([^)]*)\)", rb"/ModDate \(([^)]*)\)"):
        match = re.search(stamp, payload)
        assert match is not None
        assert not match.group(1).startswith(b"D:2026"), match.group(1)
    assert PdfReader(io.BytesIO(payload)).metadata is not None


# ── the normalised file is still the file ───────────────────────────────────


@pytest.mark.asyncio
async def test_a_normalised_workbook_reloads_with_its_data(tmp_path: Path) -> None:
    _, payload = await _publish(tmp_path, create_xlsx, sheets=SHEETS)

    workbook = load_workbook(io.BytesIO(payload))
    sheet = workbook["Summary"]

    assert sheet["A1"].value == "metric"
    assert sheet["B2"].value == 42
    assert workbook.properties.created is not None, "the property is present, just pinned"


@pytest.mark.asyncio
async def test_a_normalised_deck_reloads(tmp_path: Path) -> None:
    _, payload = await _publish(tmp_path, create_pptx, slides=SLIDES)

    deck = Presentation(io.BytesIO(payload))

    assert deck.slides[0].shapes.title.text == "Launch Readiness"


@pytest.mark.asyncio
async def test_an_invariant_pdf_still_reads(tmp_path: Path) -> None:
    _, payload = await _publish(tmp_path, create_pdf_report, title="Report", body="hello")

    reader = PdfReader(io.BytesIO(payload))

    assert "hello" in reader.pages[0].extract_text()


def test_core_xml_is_still_well_formed_after_the_rewrite() -> None:
    """The rewrite touches only the text node; the element, its attributes and
    every namespace declaration are whatever the writer emitted."""
    from openpyxl import Workbook

    out = io.BytesIO()
    Workbook().save(out)
    with zipfile.ZipFile(io.BytesIO(out.getvalue())) as archive:
        before = archive.read("docProps/core.xml")
    with zipfile.ZipFile(io.BytesIO(_normalize_zip_timestamps(out.getvalue()))) as archive:
        after = archive.read("docProps/core.xml")

    ElementTree.fromstring(after)  # parses
    strip = re.compile(rb"(<dcterms:(?:created|modified)\b[^>]*>)[^<]*")
    assert strip.sub(rb"\1", before) == strip.sub(rb"\1", after), "only the two values changed"


# ── the normaliser on its own ───────────────────────────────────────────────


def test_the_normaliser_is_idempotent() -> None:
    from openpyxl import Workbook

    out = io.BytesIO()
    Workbook().save(out)
    once = _normalize_zip_timestamps(out.getvalue())

    assert _normalize_zip_timestamps(once) == once


def test_the_normaliser_preserves_every_member_and_its_content() -> None:
    from openpyxl import Workbook

    out = io.BytesIO()
    workbook = Workbook()
    workbook.active.append(["kept", "content"])
    workbook.save(out)
    original, normalised = out.getvalue(), _normalize_zip_timestamps(out.getvalue())

    with zipfile.ZipFile(io.BytesIO(original)) as a, zipfile.ZipFile(io.BytesIO(normalised)) as b:
        assert a.namelist() == b.namelist()
        for name in a.namelist():
            if name != "docProps/core.xml":
                assert a.read(name) == b.read(name), name


def test_a_package_without_core_properties_passes_through() -> None:
    """Not every zip the normaliser might see is an OOXML package."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("a.txt", "hello")

    normalised = _normalize_zip_timestamps(out.getvalue())

    with zipfile.ZipFile(io.BytesIO(normalised)) as archive:
        assert archive.read("a.txt") == b"hello"
        assert archive.getinfo("a.txt").date_time == (1980, 1, 1, 0, 0, 0)


def test_the_core_rewrite_leaves_other_dcterms_elements_alone() -> None:
    """Only ``created`` and ``modified`` are clocks."""
    core = (
        b'<cp:coreProperties xmlns:cp="x" xmlns:dcterms="http://purl.org/dc/terms/">'
        b'<dcterms:created xsi:type="dcterms:W3CDTF">2026-09-17T09:51:48Z</dcterms:created>'
        b"<dcterms:available>2026-09-17T09:51:48Z</dcterms:available>"
        b"</cp:coreProperties>"
    )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("docProps/core.xml", core)

    with zipfile.ZipFile(io.BytesIO(_normalize_zip_timestamps(out.getvalue()))) as archive:
        rewritten = archive.read("docProps/core.xml")

    assert (
        b'<dcterms:created xsi:type="dcterms:W3CDTF">1980-01-01T00:00:00Z</dcterms:created>'
        in rewritten
    )
    assert b"<dcterms:available>2026-09-17T09:51:48Z</dcterms:available>" in rewritten


# ── different input is still a different file ───────────────────────────────


@pytest.mark.asyncio
async def test_different_sheet_data_is_a_different_artifact(tmp_path: Path) -> None:
    """Determinism must not collapse distinct outputs into one."""
    first, a = await _publish(tmp_path, create_xlsx, name="m.xlsx", sheets=SHEETS)
    other = [{"name": "Summary", "rows": [["metric", "value"], ["requests", 43]]}]
    second, b = await _publish(tmp_path, create_xlsx, name="m.xlsx", sheets=other)

    assert a != b
    assert second["status"] == "published"
    assert second["artifact"]["id"] != first["artifact"]["id"]


@pytest.mark.asyncio
async def test_a_different_session_publishes_its_own_copy(tmp_path: Path) -> None:
    """Deduplication is per session; the same bytes in another session are a
    new deliverable there."""
    import json

    for session in ("session-a", "session-b"):
        ctx = _context(tmp_path, session=session)
        token = current_tool_context.set(ctx)
        try:
            payload = json.loads(await create_xlsx(name="m.xlsx", sheets=SHEETS))
        finally:
            current_tool_context.reset(token)
        assert payload["status"] == "published"
