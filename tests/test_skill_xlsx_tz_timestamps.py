"""xlsx skill — offset-aware ISO timestamps and `set_cell` coordinates.

Two regression guards for the bundled xlsx writers:

* ``set_cell`` / ``create_xlsx`` used to feed any ISO-looking string to
  ``datetime.fromisoformat`` and write the result straight into a cell. A
  ``Z`` or ``±HH:MM`` suffix parses to an *offset-aware* datetime, which
  openpyxl refuses to serialise -- the resulting ``TypeError`` escaped as a
  raw traceback and aborted ``wb.save()`` half-written, leaving a corrupt
  3-entry zip. When ``--out`` pointed at the input file (the in-place edit)
  that destroyed the caller's only copy.
* ``set_cell`` coordinates reached ``Worksheet.cell`` unvalidated: ``0`` /
  negatives / unparseable raised raw tracebacks, and a fractional one was
  silently floored to a neighbouring cell.

Both are refused input and belong on the ``error:`` / exit 2 (or
skip-and-don't-count) paths the ops-file contract already promises.
"""

from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path
from typing import Any

import pytest

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "src" / "agentos" / "skills" / "bundled" / "xlsx" / "scripts"


def _import_scripts() -> tuple[Any, Any]:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_xlsx  # type: ignore[import-not-found]
        import edit_xlsx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)
    return create_xlsx, edit_xlsx


def _make_src(tmp_path: Path) -> Path:
    create_xlsx, _ = _import_scripts()
    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["keep me"]]}]}).save(str(src))
    return src


def _run_cli(
    edit_xlsx: Any,
    monkeypatch: pytest.MonkeyPatch,
    src: Path,
    out: Path,
    ops: list[dict[str, Any]],
    tmp_path: Path,
) -> int:
    """Run ``edit_xlsx.main()`` in-process and return its exit code."""
    ops_path = tmp_path / "ops.json"
    ops_path.write_text(json.dumps(ops), encoding="utf-8")
    monkeypatch.setattr(sys, "argv", ["edit_xlsx.py", str(src), str(ops_path), "--out", str(out)])
    return edit_xlsx.main()


def _cell(path: Path, row: int = 1, col: int = 1, sheet: str = "S") -> Any:
    from openpyxl import load_workbook

    return load_workbook(str(path))[sheet].cell(row=row, column=col).value


def _zip_entries(path: Path) -> set[str]:
    with zipfile.ZipFile(str(path)) as zf:
        return set(zf.namelist())


@pytest.mark.parametrize("value", ["2026-01-02T03:04:05Z", "2026-01-02T03:04:05+00:00"])
def test_set_cell_writes_an_offset_aware_timestamp_as_wall_clock(
    value: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """A ``Z`` / offset suffix must not abort ``wb.save`` or drop the write.

    Regression: the parsed tz-aware datetime reached ``wb.save`` and openpyxl
    refused it (``TypeError: Excel does not support timezones in datetimes``),
    which escaped as a traceback and truncated the destination zip.
    """
    _, edit_xlsx = _import_scripts()
    src = _make_src(tmp_path)
    out = tmp_path / "out.xlsx"

    rc = _run_cli(
        edit_xlsx,
        monkeypatch,
        src,
        out,
        [{"op": "set_cell", "sheet": "S", "row": 1, "col": 1, "value": value}],
        tmp_path,
    )

    capsys.readouterr()
    assert rc == 0
    # Saved as a real workbook (the corrupt 3-entry zip lacked [Content_Types]).
    assert "[Content_Types].xml" in _zip_entries(out)
    written = _cell(out)
    assert written is not None
    # The wall-clock time is kept; the offset is dropped (naive datetime).
    assert (written.year, written.month, written.day, written.hour, written.minute) == (
        2026,
        1,
        2,
        3,
        4,
    )
    assert written.tzinfo is None


def test_an_in_place_edit_survives_a_writer_refusal(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """``--out`` pointing at the input must never destroy the caller's copy.

    Regression: ``wb.save`` truncated the destination up front, so a save that
    failed partway left a corrupt 3-entry zip in place of the input -- the
    caller's only copy of the workbook. The save is now written to a sibling
    temp file and renamed over the destination only after it succeeds.
    """
    _, edit_xlsx = _import_scripts()
    src = _make_src(tmp_path)
    before = src.read_bytes()

    ops = [{"op": "set_cell", "sheet": "S", "row": 1, "col": 1, "value": "x"}]
    ops_path = tmp_path / "ops.json"
    ops_path.write_text(json.dumps(ops), encoding="utf-8")

    # Make save() fail partway through -- exactly the failure (an
    # unserialisable value) that used to truncate the destination mid-save.
    def _boom(self: Any, filename: Any) -> None:
        raise TypeError("Excel does not support timezones in datetimes")

    monkeypatch.setattr(edit_xlsx, "load_workbook", lambda **kw: _FakeBook(_boom))
    monkeypatch.setattr(sys, "argv", ["edit_xlsx.py", str(src), str(ops_path), "--out", str(src)])

    rc = edit_xlsx.main()
    capsys.readouterr()

    # Refused input is ``error:`` / exit 2 -- never a traceback.
    assert rc == 2
    # And the in-place destination is byte-for-byte what it was.
    assert src.read_bytes() == before
    assert "[Content_Types].xml" in _zip_entries(src)


class _FakeCell:
    def __init__(self) -> None:
        self.value: Any = None
        self.data_type: str = "s"
        self.quotePrefix: bool = False


class _FakeSheet:
    title = "S"

    def __init__(self) -> None:
        self._cells: dict[tuple[int, int], _FakeCell] = {}

    def cell(self, row: int, column: int) -> _FakeCell:
        return self._cells.setdefault((row, column), _FakeCell())


class _FakeBook:
    """Enough of a Workbook for ``apply_ops`` + ``_save_atomic`` to run."""

    sheetnames = ["S"]

    def __init__(self, save: Any) -> None:
        self._sheet = _FakeSheet()
        self._save = save

    def __getitem__(self, name: str) -> _FakeSheet:
        return self._sheet

    def save(self, filename: Any) -> None:
        self._save(filename)


@pytest.mark.parametrize(
    ("coord", "value"),
    [
        pytest.param({"row": 0}, 0, id="zero"),
        pytest.param({"row": -5}, -5, id="negative"),
        pytest.param({"row": "abc"}, "abc", id="unparseable"),
        pytest.param({"row": 1.7}, 1.7, id="fractional"),
        pytest.param({"row": True}, True, id="bool"),
    ],
)
def test_set_cell_refuses_an_unusable_coordinate_without_failing_the_batch(
    coord: dict[str, Any],
    value: Any,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """``0`` / negative / unparseable / fractional / bool coordinates are bad
    input: the op is skipped and uncounted, never a raw traceback -- and a
    fractional one must not silently floor to a neighbouring cell.

    Regression: ``row=0``/``-5``/``"abc"`` raised ``ValueError`` from
    ``Worksheet.cell`` / ``int()``, and ``row=1.7`` was quietly truncated to
    ``1``, reporting ``{"applied": 1}`` while writing a different cell than
    the op named.
    """
    _, edit_xlsx = _import_scripts()
    src = _make_src(tmp_path)
    out = tmp_path / "out.xlsx"

    rc = _run_cli(
        edit_xlsx,
        monkeypatch,
        src,
        out,
        [
            {"op": "set_cell", "sheet": "S", "col": 1, "value": "should land", **coord},
            {"op": "set_cell", "sheet": "S", "row": 1, "col": 1, "value": "should land"},
        ],
        tmp_path,
    )

    report = json.loads(capsys.readouterr().out.strip())
    assert rc == 0
    # Only the good op applied -- the refused one was skipped and uncounted.
    assert report == {"applied": 1}
    assert _cell(out) == "should land"


def test_set_cell_accepts_a_whole_float_coordinate(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """``1.0`` names row 1 and must keep working -- only a non-integral float
    is refused."""
    _, edit_xlsx = _import_scripts()
    src = _make_src(tmp_path)
    out = tmp_path / "out.xlsx"

    rc = _run_cli(
        edit_xlsx,
        monkeypatch,
        src,
        out,
        [{"op": "set_cell", "sheet": "S", "row": 1.0, "col": 1, "value": "ok"}],
        tmp_path,
    )

    report = json.loads(capsys.readouterr().out.strip())
    assert rc == 0
    assert report == {"applied": 1}
    assert _cell(out) == "ok"


def test_create_xlsx_writes_an_offset_aware_timestamp_as_wall_clock(
    tmp_path: Path,
) -> None:
    """``create_xlsx`` shares the coercion and must not abort ``wb.save``."""
    create_xlsx, _ = _import_scripts()
    out = tmp_path / "created.xlsx"
    wb = create_xlsx.build({"sheets": [{"name": "S", "rows": [["2026-01-02T03:04:05Z"]]}]})
    create_xlsx._save_atomic(wb, out)

    from openpyxl import load_workbook

    written = load_workbook(str(out))["S"]["A1"].value
    assert written.tzinfo is None
    assert (written.year, written.month, written.day, written.hour, written.minute) == (
        2026,
        1,
        2,
        3,
        4,
    )


def test_create_xlsx_in_place_target_survives_a_writer_refusal(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """``create_xlsx._save_atomic`` must leave the destination alone on error.

    Regression: a save that failed partway left a corrupt 3-entry zip in
    place of the destination -- for an in-place ``--out`` that was the
    caller's only copy of the workbook. The save is written to a sibling
    temp file and renamed over the destination only after it succeeds.
    """
    create_xlsx, _ = _import_scripts()
    dest = tmp_path / "book.xlsx"
    # Seed the destination with a valid workbook the caller cares about.
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["keep me"]]}]}).save(str(dest))
    before = dest.read_bytes()

    class _Boom:
        sheetnames = ["S"]

        def save(self, filename: Any) -> None:
            raise TypeError("Excel does not support timezones in datetimes")

    with pytest.raises(TypeError):
        create_xlsx._save_atomic(_Boom(), dest)
    capsys.readouterr()

    # The in-place destination is byte-for-byte what it was.
    assert dest.read_bytes() == before
    assert "[Content_Types].xml" in _zip_entries(dest)
    # And no temp file is left behind.
    assert not list(tmp_path.glob("*.tmp"))


def test_save_atomic_leaves_no_temp_file_after_a_clean_save(tmp_path: Path) -> None:
    """``_save_atomic`` writes through ``<target>.xlsx.tmp`` and renames it in.

    Regression: a leftover ``*.xlsx.tmp`` would pollute the caller's
    directory after a *successful* save. The rename is the entire point --
    after a clean save there must be only the destination.
    """
    create_xlsx, _ = _import_scripts()
    dest = tmp_path / "clean.xlsx"
    wb = create_xlsx.build({"sheets": [{"name": "S", "rows": [["ok"]]}]})
    create_xlsx._save_atomic(wb, dest)

    assert dest.exists()
    assert _cell(dest) == "ok"
    assert not list(tmp_path.glob("*.tmp")), "temp file should be renamed in, not left behind"
