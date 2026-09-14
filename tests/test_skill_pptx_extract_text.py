"""pptx skill extract_text unit tests."""

from __future__ import annotations

import io
import json
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from agentos.skills.bundled.pptx.scripts import extract_text


def test_shape_text_and_table_text_extraction(tmp_path: Path) -> None:
    """extract_text extracts paragraph text and multi-paragraph table cells cleanly."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 1. Textbox shape with paragraph text
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "Slide Headline"
    p_body = tb.text_frame.add_paragraph()
    p_body.text = "Secondary point"

    # 2. Table with multi-paragraph cell
    table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(2.5), Inches(5), Inches(2))
    table = table_shape.table
    table.cell(0, 0).text_frame.text = "Metric"
    table.cell(0, 1).text_frame.text = "Value"

    c1 = table.cell(1, 0)
    c1.text_frame.text = "Q1 Revenue"
    p_extra = c1.text_frame.add_paragraph()
    p_extra.text = "(USD)"

    c2 = table.cell(1, 1)
    c2.text_frame.text = "$100M"

    # Test helpers directly
    shape_lines = extract_text._shape_text(tb)
    assert shape_lines == ["Slide Headline", "Secondary point"]

    table_lines = extract_text._table_text(table_shape)
    assert table_lines == ["Metric | Value", "Q1 Revenue (USD) | $100M"]

    # Save presentation and test CLI entrypoint
    pptx_path = tmp_path / "deck.pptx"
    prs.save(str(pptx_path))

    # Test main with --json
    stdout_capture = io.StringIO()
    monkeypatch = pytest.MonkeyPatch()
    monkeypatch.setattr(sys, "stdout", stdout_capture)
    try:
        ret = extract_text.main([str(pptx_path), "--json"])
        assert ret == 0
    finally:
        monkeypatch.undo()

    payload = json.loads(stdout_capture.getvalue())
    assert len(payload) == 1
    slide_data = payload[0]
    assert slide_data["slide"] == 1
    assert "Slide Headline" in slide_data["text"]
    assert "Secondary point" in slide_data["text"]
    assert "Metric | Value" in slide_data["text"]
    assert "Q1 Revenue (USD) | $100M" in slide_data["text"]

def test_nested_group_shapes_are_walked_recursively() -> None:
    """Text and tables nested two+ group levels deep must not be dropped."""

    class _Para:
        def __init__(self, text: str) -> None:
            self.text = text

    class _Cell:
        def __init__(self, text: str) -> None:
            self.text_frame = type("Tf", (), {"paragraphs": [_Para(text)]})()

    class _Row:
        def __init__(self, *texts: str) -> None:
            self.cells = [_Cell(t) for t in texts]

    class _Table:
        def __init__(self, *rows: _Row) -> None:
            self.rows = list(rows)

    class _Frame:
        def __init__(self, paragraphs: list[_Para]) -> None:
            self.paragraphs = paragraphs

    class _Shape:
        def __init__(
            self,
            *,
            frame: _Frame | None = None,
            table: _Table | None = None,
            children: tuple = (),
        ) -> None:
            self.has_text_frame = frame is not None
            self.text_frame = frame
            self.has_table = table is not None
            self.table = table
            self.shape_type = len(children) > 0
            self.shapes = list(children) if children else None

    leaf = _Shape(frame=_Frame([_Para("deepest leaf")]))
    table = _Shape(table=_Table(_Row("L1C1", "L1C2")))
    inner_group = _Shape(children=(leaf, table))
    outer_group = _Shape(children=(inner_group,))
    slide = type("Slide", (), {"shapes": [outer_group]})()

    lines = extract_text._slide_text(slide)
    assert "deepest leaf" in lines
    assert "L1C1 | L1C2" in lines

