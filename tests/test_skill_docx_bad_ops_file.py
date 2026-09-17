"""Issue #2463: ``edit_docx.py`` reported success for an ops file it could not use.

``ops = raw if isinstance(raw, list) else []`` turned a single operation
object -- the common mistake instead of a one-element list -- and any other
non-list JSON into "no operations", then opened the document, applied
nothing, wrote an unedited copy over ``--out`` and printed ``{"applied": 0}``
with exit 0. Invalid JSON, or the UTF-16 PowerShell's ``Out-File`` writes by
default, escaped as a traceback with exit 1.

A second silence sat behind the first: a well-formed list whose operations
match nothing -- a typo in ``op``, a ``find`` that is not in the document, a
``para`` index off the end -- produced the same ``{"applied": 0}`` and the
same unedited copy, with no indication which operation was wrong or why.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from types import ModuleType

import pytest
from docx import Document

SCRIPTS = Path(__file__).resolve().parent.parent / "src/agentos/skills/bundled/docx/scripts"


@pytest.fixture(scope="module")
def edit_docx() -> ModuleType:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import edit_docx as module  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)
    return module


@pytest.fixture
def source(tmp_path: Path) -> Path:
    """A document with one placeholder, and a run to address by index."""
    document = Document()
    document.add_paragraph("Hello {{NAME}}")
    path = tmp_path / "doc.docx"
    document.save(str(path))
    return path


def _run(
    edit_docx: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    source: Path,
    ops_path: Path,
    out: Path,
) -> int:
    monkeypatch.setattr(
        sys, "argv", ["edit_docx.py", str(source), str(ops_path), "--out", str(out)]
    )
    return edit_docx.main()


def _ops(tmp_path: Path, body: str | bytes, name: str = "ops.json") -> Path:
    path = tmp_path / name
    if isinstance(body, bytes):
        path.write_bytes(body)
    else:
        path.write_text(body, encoding="utf-8")
    return path


def _text(path: Path) -> str:
    return "\n".join(p.text for p in Document(str(path)).paragraphs)


# ── the issue: a file that is not a list ────────────────────────────────────


def test_the_issues_single_object_is_refused_and_nothing_is_written(
    edit_docx: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    tmp_path: Path,
    source: Path,
) -> None:
    ops = _ops(tmp_path, '{"op": "replace_text", "find": "{{NAME}}", "with": "World"}')
    out = tmp_path / "out.docx"

    code = _run(edit_docx, monkeypatch, source, ops, out)

    captured = capsys.readouterr()
    assert code == 2
    assert captured.out == ""
    assert captured.err.startswith(f"error: ops {ops} must be a JSON list of operations, got dict")
    assert not out.exists()


@pytest.mark.parametrize(
    ("body", "type_name"),
    [
        ('{"op": "replace_text"}', "dict"),
        ('{"ops": [{"op": "replace_text"}]}', "dict"),
        ("null", "NoneType"),
        ('"replace"', "str"),
        ("42", "int"),
        ("4.2", "float"),
        ("true", "bool"),
        ("false", "bool"),
    ],
)
def test_every_non_list_shape_is_refused_with_its_type_named(
    edit_docx: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    tmp_path: Path,
    source: Path,
    body: str,
    type_name: str,
) -> None:
    ops = _ops(tmp_path, body)
    out = tmp_path / "out.docx"

    code = _run(edit_docx, monkeypatch, source, ops, out)

    captured = capsys.readouterr()
    assert code == 2
    assert captured.out == ""
    assert f"got {type_name}" in captured.err
    assert not out.exists()


def test_an_existing_output_is_not_overwritten_by_a_refusal(
    edit_docx: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    source: Path,
) -> None:
    """The damaging case from the issue: --out pointing at a real document."""
    out = tmp_path / "out.docx"
    good = Document()
    good.add_paragraph("previous good content")
    good.save(str(out))
    before = out.read_bytes()

    code = _run(edit_docx, monkeypatch, source, _ops(tmp_path, "null"), out)

    assert code == 2
    assert out.read_bytes() == before


def test_no_output_directory_is_created_by_a_refusal(
    edit_docx: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    source: Path,
) -> None:
    """``main`` mkdirs the parent before writing; the refusal has to come first."""
    out = tmp_path / "new" / "dir" / "out.docx"

    code = _run(edit_docx, monkeypatch, source, _ops(tmp_path, "42"), out)

    assert code == 2
    assert not out.parent.exists()


# ── the issue: a file that is not JSON, or not UTF-8 ────────────────────────


@pytest.mark.parametrize(
    "body",
    [
        pytest.param(b'[{"op": "replace_text",', id="truncated"),
        pytest.param(b"", id="empty"),
        pytest.param(b"   \n", id="whitespace"),
        pytest.param(b'[{"op": "replace_text"},]', id="trailing-comma"),
        pytest.param(b"[{'op': 'replace_text'}]", id="single-quotes"),
        pytest.param(b"replace_text", id="bare-word"),
        pytest.param(b"\xff\xfe" + '[{"op": "x"}]'.encode("utf-16-le"), id="utf-16le-bom"),
        pytest.param('[{"op": "x"}]'.encode("utf-16-be"), id="utf-16be"),
        pytest.param(b"[\xff\xfe]", id="invalid-utf-8"),
    ],
)
def test_an_unparseable_file_is_refused_with_no_traceback(
    edit_docx: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    tmp_path: Path,
    source: Path,
    body: bytes,
) -> None:
    ops = _ops(tmp_path, body)
    out = tmp_path / "out.docx"

    code = _run(edit_docx, monkeypatch, source, ops, out)

    captured = capsys.readouterr()
    assert code == 2
    assert captured.out == ""
    assert captured.err.startswith(f"error: ops {ops} is not")
    assert "Traceback" not in captured.err
    assert not out.exists()


def test_a_utf16_file_gets_a_message_that_names_the_cause(
    edit_docx: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    tmp_path: Path,
    source: Path,
) -> None:
    """PowerShell 5.1's Out-File writes UTF-16LE with a BOM by default; the
    person who hit this needs to be told to re-save, not shown a codec error."""
    ops = _ops(tmp_path, '[{"op": "replace_text", "find": "a", "with": "b"}]'.encode("utf-16"))

    _run(edit_docx, monkeypatch, source, ops, tmp_path / "out.docx")

    err = capsys.readouterr().err
    assert "not UTF-8" in err
    assert "UTF-16" in err and "Out-File" in err


def test_the_input_is_never_modified_by_a_refusal(
    edit_docx: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    source: Path,
) -> None:
    before = source.read_bytes()

    _run(edit_docx, monkeypatch, source, _ops(tmp_path, b"{"), tmp_path / "out.docx")

    assert source.read_bytes() == before


def test_the_pre_existing_not_found_checks_still_fire_first(
    edit_docx: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    tmp_path: Path,
    source: Path,
) -> None:
    code = _run(edit_docx, monkeypatch, source, tmp_path / "missing.json", tmp_path / "o.docx")
    assert code == 2
    assert "not found" in capsys.readouterr().err

    code = _run(
        edit_docx, monkeypatch, tmp_path / "missing.docx", _ops(tmp_path, "[]"), tmp_path / "o.docx"
    )
    assert code == 2
    assert "input" in capsys.readouterr().err


# ── a well-formed list is still accepted ────────────────────────────────────


def test_a_valid_list_applies_and_reports_as_before(
    edit_docx: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    tmp_path: Path,
    source: Path,
) -> None:
    ops = _ops(tmp_path, '[{"op": "replace_text", "find": "{{NAME}}", "with": "World"}]')
    out = tmp_path / "out.docx"

    code = _run(edit_docx, monkeypatch, source, ops, out)

    captured = capsys.readouterr()
    assert code == 0
    assert json.loads(captured.out) == {"applied": 1}
    assert captured.err == ""
    assert _text(out) == "Hello World"


def test_an_empty_list_is_a_valid_request_for_no_edits(
    edit_docx: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    tmp_path: Path,
    source: Path,
) -> None:
    """Unlike a non-list, ``[]`` is well-formed; it copies the document and
    says nothing on stderr, because nothing was asked for."""
    out = tmp_path / "out.docx"

    code = _run(edit_docx, monkeypatch, source, _ops(tmp_path, "[]"), out)

    captured = capsys.readouterr()
    assert code == 0
    assert json.loads(captured.out) == {"applied": 0}
    assert captured.err == ""
    assert _text(out) == "Hello {{NAME}}"


# ── the second silence: a list whose operations apply nothing ───────────────


@pytest.mark.parametrize(
    ("op", "reason"),
    [
        ({"op": "replace_txt", "find": "{{NAME}}", "with": "x"}, "unknown op 'replace_txt'"),
        ({"find": "{{NAME}}", "with": "x"}, "has no 'op' key"),
        ({"op": "replace_text", "find": "{{MISSING}}", "with": "x"}, "found '{{MISSING}}' nowhere"),
        ({"op": "replace_text", "with": "x"}, "empty 'find'"),
        ({"op": "replace_text", "find": "", "with": "x"}, "empty 'find'"),
        ({"op": "replace_run", "run": 0, "text": "x"}, "no 'para' index"),
        ({"op": "replace_run", "para": "first", "text": "x"}, "non-integer index"),
        ({"op": "replace_run", "para": 7, "text": "x"}, "names paragraph 7; the document has 1"),
        ({"op": "replace_run", "para": -1, "text": "x"}, "names paragraph -1"),
        ({"op": "replace_run", "para": 0, "run": 5, "text": "x"}, "names run 5; paragraph 0 has 1"),
    ],
)
def test_each_skipped_operation_is_explained_on_stderr(
    edit_docx: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    tmp_path: Path,
    source: Path,
    op: dict[str, object],
    reason: str,
) -> None:
    out = tmp_path / "out.docx"

    code = _run(edit_docx, monkeypatch, source, _ops(tmp_path, json.dumps([op])), out)

    captured = capsys.readouterr()
    assert code == 0, "a skipped op is a warning, not a failure"
    assert json.loads(captured.out) == {"applied": 0}, "the stdout contract is unchanged"
    assert "warning: op 0: " in captured.err
    assert reason in captured.err
    assert "none of the 1 operations applied" in captured.err
    assert f"{out} is a copy of the input" in captured.err


def test_a_non_object_element_is_explained_and_does_not_cost_the_batch(
    edit_docx: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    tmp_path: Path,
    source: Path,
) -> None:
    ops = _ops(
        tmp_path,
        json.dumps(["oops", {"op": "replace_text", "find": "{{NAME}}", "with": "World"}]),
    )
    out = tmp_path / "out.docx"

    code = _run(edit_docx, monkeypatch, source, ops, out)

    captured = capsys.readouterr()
    assert code == 0
    assert json.loads(captured.out) == {"applied": 1}
    assert "warning: op 0: is not an object (got str)" in captured.err
    assert "none of the" not in captured.err, "one applied, so the batch did not fail"
    assert _text(out) == "Hello World"


def test_a_partly_applied_batch_warns_about_the_right_op_only(
    edit_docx: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    tmp_path: Path,
    source: Path,
) -> None:
    ops = _ops(
        tmp_path,
        json.dumps(
            [
                {"op": "replace_text", "find": "{{NAME}}", "with": "World"},
                {"op": "replace_text", "find": "{{NOPE}}", "with": "x"},
            ]
        ),
    )

    _run(edit_docx, monkeypatch, source, ops, tmp_path / "out.docx")

    err = capsys.readouterr().err
    assert "warning: op 1: " in err
    assert "warning: op 0: " not in err


def test_replace_text_still_counts_one_per_paragraph_changed(
    edit_docx: ModuleType, tmp_path: Path
) -> None:
    """The ``applied`` semantics other callers rely on are unchanged."""
    document = Document()
    document.add_paragraph("{{X}} one")
    document.add_paragraph("{{X}} two")
    document.add_paragraph("no placeholder")

    applied = edit_docx.apply_ops(document, [{"op": "replace_text", "find": "{{X}}", "with": "Y"}])

    assert applied == 2


# ── apply_ops keeps its two-argument form ───────────────────────────────────


def test_apply_ops_without_the_skipped_list_behaves_exactly_as_before(
    edit_docx: ModuleType,
) -> None:
    """``tests/test_skill_docx.py`` calls ``apply_ops(doc, ops)`` directly in
    a dozen places; the new parameter is optional and changes no result."""
    document = Document()
    document.add_paragraph("Hello {{NAME}}")
    ops = [
        {"op": "replace_text", "find": "{{NAME}}", "with": "World"},
        {"op": "replace_txt", "find": "x", "with": "y"},
        "not an op",
    ]

    assert edit_docx.apply_ops(document, ops) == 1
    assert document.paragraphs[0].text == "Hello World"


def test_apply_ops_reports_one_reason_per_skipped_op_in_order(edit_docx: ModuleType) -> None:
    document = Document()
    document.add_paragraph("Hello {{NAME}}")
    skipped: list[str] = []
    ops = [
        {"op": "replace_txt"},
        {"op": "replace_text", "find": "{{NAME}}", "with": "World"},
        {"op": "replace_run", "para": 9, "text": "x"},
    ]

    applied = edit_docx.apply_ops(document, ops, skipped)

    assert applied == 1
    assert [s.split(":")[0] for s in skipped] == ["op 0", "op 2"]


def test_a_negative_run_index_is_reported_not_wrapped(edit_docx: ModuleType) -> None:
    """A negative index must never wrap round to the last run and edit
    something the op did not name."""
    document = Document()
    document.add_paragraph("keep me")
    skipped: list[str] = []

    applied = edit_docx.apply_ops(
        document, [{"op": "replace_run", "para": 0, "run": -1, "text": "x"}], skipped
    )

    assert applied == 0
    assert document.paragraphs[0].text == "keep me"
    assert "names run -1" in skipped[0]


# ── the skill doc states the contract ───────────────────────────────────────


def test_the_skill_doc_states_the_list_contract_and_the_refusal() -> None:
    doc = (SCRIPTS.parent / "SKILL.md").read_text(encoding="utf-8")

    assert "must be a JSON **list**" in doc
    assert "exit code 2" in doc
    assert "warning: op" in doc
