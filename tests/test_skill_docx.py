"""docx skill — load, eligibility, and create→inspect round-trip."""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

from agentos.skills.eligibility import EligibilityContext, check_eligibility
from agentos.skills.loader import SkillLoader

ROOT = Path(__file__).resolve().parents[1]
BUNDLED = ROOT / "src" / "agentos" / "skills" / "bundled"
DOCX_DIR = BUNDLED / "docx"
SCRIPTS = DOCX_DIR / "scripts"


def _spec_to_loader() -> object:
    return SkillLoader(bundled_dir=BUNDLED).get_by_name("docx")


def test_skill_loads() -> None:
    spec = _spec_to_loader()
    assert spec is not None
    assert spec.name == "docx"
    assert spec.metadata is not None
    assert spec.provenance.origin == "clawhub-mit0"
    assert spec.provenance.license == "MIT-0"


def test_eligibility_with_python_present(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "agentos.skills.eligibility.shutil.which",
        lambda name: "/usr/bin/python3" if name in {"python", "python3"} else None,
    )
    spec = _spec_to_loader()
    assert spec is not None
    assert check_eligibility(spec, EligibilityContext.auto())


def test_eligibility_without_python(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "agentos.skills.eligibility.shutil.which",
        lambda name: None,
    )
    spec = _spec_to_loader()
    assert spec is not None
    assert not check_eligibility(spec, EligibilityContext.auto())


def test_create_then_inspect_round_trip(tmp_path: Path) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_docx  # type: ignore[import-not-found]
        import inspect_docx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    spec = {
        "metadata": {"title": "Round-trip", "author": "Tester"},
        "body": [
            {"kind": "heading", "level": 1, "text": "Hello"},
            {"kind": "paragraph", "text": "World."},
            {"kind": "table", "rows": [["A", "B"], ["1", "2"]]},
        ],
    }
    out_path = tmp_path / "out.docx"
    doc = create_docx.build(spec)
    doc.save(str(out_path))
    assert out_path.exists()

    inspected = inspect_docx.inspect(out_path)
    assert inspected["sections"] >= 1
    texts = [p["text"] for p in inspected["paragraphs"]]
    assert "Hello" in texts
    assert "World." in texts
    assert inspected["tables"] and inspected["tables"][0][0] == ["A", "B"]
    assert inspected["has_tracked_changes"] is False


def test_edit_replace_text(tmp_path: Path) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_docx  # type: ignore[import-not-found]
        import edit_docx  # type: ignore[import-not-found]
        import inspect_docx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    src = tmp_path / "src.docx"
    create_docx.build({"body": [{"kind": "paragraph", "text": "Hello {{NAME}}, welcome."}]}).save(
        str(src)
    )

    from docx import Document

    doc = Document(str(src))
    ops = [{"op": "replace_text", "find": "{{NAME}}", "with": "Wei"}]
    edit_docx.apply_ops(doc, ops)
    out = tmp_path / "out.docx"
    doc.save(str(out))

    inspected = inspect_docx.inspect(out)
    text = " ".join(p["text"] for p in inspected["paragraphs"])
    assert "{{NAME}}" not in text
    assert "Wei" in text


def test_inspect_cli_outputs_json(tmp_path: Path) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_docx  # type: ignore[import-not-found]
        import inspect_docx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    src = tmp_path / "src.docx"
    create_docx.build({"body": [{"kind": "paragraph", "text": "x"}]}).save(str(src))

    payload = inspect_docx.inspect(src)
    encoded = json.dumps(payload, ensure_ascii=False)
    assert "paragraphs" in encoded
    assert "tables" in encoded


def test_inspect_docx_creates_parent_directory(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_docx  # type: ignore[import-not-found]
        import inspect_docx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    src = tmp_path / "src.docx"
    create_docx.build({"body": [{"kind": "paragraph", "text": "x"}]}).save(str(src))

    out = tmp_path / "nested" / "dir" / "out.json"
    monkeypatch.setattr(sys, "argv", ["inspect_docx.py", str(src), "--out", str(out)])
    assert inspect_docx.main() == 0
    assert out.is_file()


def _edit_docx_module() -> object:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import edit_docx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)
    return edit_docx


def _paragraph(runs: list[tuple[str, bool]]) -> object:
    """Build a one-paragraph document whose runs carry the given bold flags."""
    from docx import Document

    document = Document()
    paragraph = document.add_paragraph()
    for text, bold in runs:
        run = paragraph.add_run(text)
        if bold:
            run.bold = True
    return paragraph


def test_replace_text_keeps_the_formatting_of_untouched_runs() -> None:
    """A run the match never touched must survive byte-identical.

    Runs are where Word stores character formatting, so collapsing a paragraph
    into `runs[0]` gave every character run 0's formatting and left the rest as
    empty shells -- bold, italic, font and colour discarded for the whole
    paragraph even though one word changed.
    """
    edit_docx = _edit_docx_module()
    paragraph = _paragraph([("Hello ", False), ("world", True), (" and more", False)])

    assert edit_docx._replace_text_in_paragraph(paragraph, "Hello", "Hi") is True

    assert [(run.text, run.bold) for run in paragraph.runs] == [
        ("Hi ", None),
        ("world", True),
        (" and more", None),
    ]


def test_replace_text_across_a_run_boundary_keeps_the_second_run() -> None:
    """A `find` spanning runs still leaves the surrounding text in place.

    The replacement lands in the run owning the match's first character; the
    rest of the second run keeps its own text and formatting.
    """
    edit_docx = _edit_docx_module()
    paragraph = _paragraph([("Hel", False), ("lo world", True)])

    assert edit_docx._replace_text_in_paragraph(paragraph, "Hello", "Hi") is True

    assert [(run.text, run.bold) for run in paragraph.runs] == [("Hi", None), (" world", True)]


def test_replace_text_handles_several_matches_without_moving_text() -> None:
    """Each match is replaced where it starts, so runs keep their own share."""
    edit_docx = _edit_docx_module()
    paragraph = _paragraph([("aXa", False), ("Xa", True)])

    assert edit_docx._replace_text_in_paragraph(paragraph, "X", "-") is True

    assert [(run.text, run.bold) for run in paragraph.runs] == [("a-a", None), ("-a", True)]


@pytest.mark.parametrize(
    ("runs", "find", "replacement"),
    [
        ([("Hello ", False), ("world", True), (" and more", False)], "Hello", "Hi"),
        ([("Hel", False), ("lo world", True)], "Hello", "Hi"),
        ([("a", False), ("b", True), ("c", False)], "abc", "X"),
        ([("x{{N}}y", False)], "{{N}}", "Wei"),
        ([("aXa", False), ("Xa", True)], "X", "-"),
        ([("keep ", False), ("me", True)], "me", ""),
        ([("aa", False), ("aa", True)], "aa", "b"),
    ],
)
def test_replace_text_matches_str_replace_on_the_joined_text(
    runs: list[tuple[str, bool]], find: str, replacement: str
) -> None:
    """Whatever the run layout, the resulting text is plain `str.replace`.

    The old code already got the text right -- it was the run layout that was
    wrong -- so this pins the half that must not change while the fix moves
    characters back into their own runs.
    """
    edit_docx = _edit_docx_module()
    paragraph = _paragraph(runs)
    original = "".join(text for text, _ in runs)

    edit_docx._replace_text_in_paragraph(paragraph, find, replacement)

    assert "".join(run.text for run in paragraph.runs) == original.replace(find, replacement)


def test_replace_text_reports_false_when_the_needle_is_absent() -> None:
    """No match means no edit and no reported change."""
    edit_docx = _edit_docx_module()
    paragraph = _paragraph([("Hello ", False), ("world", True)])

    assert edit_docx._replace_text_in_paragraph(paragraph, "absent", "x") is False
    assert [(run.text, run.bold) for run in paragraph.runs] == [("Hello ", None), ("world", True)]


def test_replace_run_still_touches_only_its_own_run() -> None:
    """`replace_run` was never affected; keep it that way."""
    edit_docx = _edit_docx_module()
    paragraph = _paragraph([("Hello ", False), ("world", True)])

    edit_docx._replace_run(paragraph, 1, "there")

    assert [(run.text, run.bold) for run in paragraph.runs] == [("Hello ", None), ("there", True)]


def test_replace_text_reaches_table_cells(tmp_path: Path) -> None:
    """Placeholders in contracts and invoices usually live inside tables.

    `apply_ops` only walked `doc.paragraphs`, which python-docx limits to the
    body, so a `{{CLIENT}}` in a table cell was never replaced and the op
    reported zero applications.
    """
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    doc.add_paragraph("Agreement Header")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Client:"
    table.cell(0, 1).text = "{{CLIENT}}"

    applied = edit_docx.apply_ops(
        doc, [{"op": "replace_text", "find": "{{CLIENT}}", "with": "Acme Corp"}]
    )

    assert applied == 1
    assert table.cell(0, 1).text == "Acme Corp"
    assert table.cell(0, 0).text == "Client:"


def test_replace_text_counts_body_and_table_paragraphs_together() -> None:
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    doc.add_paragraph("Dear {{NAME}},")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Signed: {{NAME}}"

    applied = edit_docx.apply_ops(doc, [{"op": "replace_text", "find": "{{NAME}}", "with": "Wei"}])

    assert applied == 2
    assert doc.paragraphs[0].text == "Dear Wei,"
    assert doc.tables[0].cell(0, 0).text == "Signed: Wei"


def test_replace_text_reaches_nested_tables() -> None:
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "Total: {{TOTAL}}"

    applied = edit_docx.apply_ops(
        doc, [{"op": "replace_text", "find": "{{TOTAL}}", "with": "42.00"}]
    )

    assert applied == 1
    assert inner.cell(0, 0).text == "Total: 42.00"


def test_replace_text_visits_a_merged_cell_once() -> None:
    """`row.cells` repeats a merged cell for every grid column it spans.

    Walking it once per column would apply the replacement again to text the
    first pass already rewrote; a replacement containing its own needle makes
    that visible.
    """
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    table = doc.add_table(rows=1, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 2))
    merged.text = "{{X}}"

    applied = edit_docx.apply_ops(doc, [{"op": "replace_text", "find": "{{X}}", "with": "{{X}}!"}])

    assert applied == 1
    assert merged.text == "{{X}}!"


def test_replace_text_survives_an_irregular_vertical_merge() -> None:
    """`row.cells` resolves a `vMerge=continue` cell against the row above and
    raises `ValueError` when no cell starts at that grid offset there -- a
    layout non-Word generators produce. Walking the `<w:tc>` elements directly
    never enters that path, so the whole edit (body included) still lands."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    edit_docx = _edit_docx_module()
    doc = Document()
    doc.add_paragraph("Header {{X}}")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 0).text = "Cell {{X}}"
    continue_marker = OxmlElement("w:vMerge")
    continue_marker.set(qn("w:val"), "continue")
    table.cell(1, 1)._tc.get_or_add_tcPr().append(continue_marker)

    applied = edit_docx.apply_ops(doc, [{"op": "replace_text", "find": "{{X}}", "with": "Y"}])

    assert applied == 2
    assert doc.paragraphs[0].text == "Header Y"
    assert table.cell(1, 0).text == "Cell Y"


def test_apply_ops_skips_non_dict_ops() -> None:
    """Malformed op lists are ignored, as `edit_xlsx.apply_ops` already does."""
    from docx import Document

    edit_docx = _edit_docx_module()
    doc = Document()
    doc.add_paragraph("Hello {{NAME}}")

    applied = edit_docx.apply_ops(
        doc,
        [None, "replace_text", 3, {"op": "replace_text", "find": "{{NAME}}", "with": "Wei"}],
    )

    assert applied == 1
    assert doc.paragraphs[0].text == "Hello Wei"


def _create_docx_module() -> object:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_docx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)
    return create_docx


def test_build_handles_non_dict_spec() -> None:
    """Passing non-dict or malformed specs should safely return an empty document."""
    create_docx = _create_docx_module()

    for invalid_spec in [None, [], "not-a-dict", 123]:
        doc = create_docx.build(invalid_spec)  # type: ignore[attr-defined]
        assert doc is not None
        assert len(doc.paragraphs) == 0
        assert len(doc.tables) == 0


def test_build_handles_empty_or_malformed_table_rows() -> None:
    """Zero-column tables or non-list row elements should not crash Document.add_table."""
    create_docx = _create_docx_module()

    # Empty sub-rows (ncols == 0) should be skipped without error
    doc1 = create_docx.build({"body": [{"kind": "table", "rows": [[], []]}]})  # type: ignore[attr-defined]
    assert len(doc1.tables) == 0

    # Non-list row items should be coerced safely into cells
    doc2 = create_docx.build({"body": [{"kind": "table", "rows": ["header", 42]}]})  # type: ignore[attr-defined]
    assert len(doc2.tables) == 1
    assert doc2.tables[0].rows[0].cells[0].text == "header"
    assert doc2.tables[0].rows[1].cells[0].text == "42"

    # Jagged rows should be padded up to ncols without IndexError
    doc3 = create_docx.build(  # type: ignore[attr-defined]
        {"body": [{"kind": "table", "rows": [["a", "b", "c"], ["1"]]}]}
    )
    assert len(doc3.tables) == 1
    assert len(doc3.tables[0].columns) == 3
    assert doc3.tables[0].rows[1].cells[0].text == "1"


def test_build_handles_invalid_heading_levels() -> None:
    """Out-of-range or invalid level values should be clamped safely to 0-9."""
    create_docx = _create_docx_module()

    spec = {
        "body": [
            {"kind": "heading", "text": "Clamped Min", "level": -5},
            {"kind": "heading", "text": "Clamped Max", "level": 50},
            {"kind": "heading", "text": "Non-integer", "level": "bad"},
            {"kind": "heading", "text": "None Level", "level": None},
        ]
    }
    doc = create_docx.build(spec)  # type: ignore[attr-defined]
    assert len(doc.paragraphs) == 4
    assert doc.paragraphs[0].text == "Clamped Min"
    assert doc.paragraphs[1].text == "Clamped Max"
    assert doc.paragraphs[2].text == "Non-integer"
    assert doc.paragraphs[3].text == "None Level"


def test_create_docx_cli_errors_on_malformed_spec(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """CLI exits with code 2 on missing, unparseable, or non-object JSON specs."""
    create_docx = _create_docx_module()

    non_dict_spec = tmp_path / "spec_list.json"
    non_dict_spec.write_text("[]", encoding="utf-8")
    out = tmp_path / "out.docx"

    monkeypatch.setattr(sys, "argv", ["create_docx.py", str(non_dict_spec), "--out", str(out)])
    assert create_docx.main() == 2  # type: ignore[attr-defined]
    assert not out.exists()

    invalid_json = tmp_path / "bad.json"
    invalid_json.write_text("{invalid json", encoding="utf-8")
    monkeypatch.setattr(sys, "argv", ["create_docx.py", str(invalid_json), "--out", str(out)])
    assert create_docx.main() == 2  # type: ignore[attr-defined]
